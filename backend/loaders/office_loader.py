"""
Office document loader: DOCX only
"""
import docx
import os

HEADING_STYLES = {"Heading 1": 1, "Heading 2": 2, "Heading 3": 3, "Title": 1}


def load_docx(file_path: str):
    filename = os.path.basename(file_path)
    document = docx.Document(file_path)
    elements, section, subsection = [], None, None

    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = para.style.name if para.style else ""

        if style in HEADING_STYLES:
            level = HEADING_STYLES[style]
            section, subsection = (text, None) if level == 1 else (section, text)
            elements.append({
                "type": "heading", "content": text, "page": 1,
                "metadata": {"level": level, "source_file": filename}
            })
        else:
            elements.append({
                "type": "paragraph", "content": text, "page": 1,
                "metadata": {"section": section, "subsection": subsection, "source_file": filename}
            })

    for i, table in enumerate(document.tables):
        rows = ["\t".join(c.text.strip() for c in row.cells) for row in table.rows]
        elements.append({
            "type": "table", "content": "\n".join(rows), "page": 1,
            "metadata": {"table_id": f"Table_{i+1}", "section": section, "source_file": filename}
        })

    return elements
